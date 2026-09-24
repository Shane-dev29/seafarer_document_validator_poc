import openpyxl
from typing import Optional, Dict
from core.config import EXCEL_PATH
from core.schemas import PortalEntry

# Map common demonyms/nationalities and ISO codes to country names
NATIONALITY_ALIASES: Dict[str, str] = {
    "indian": "India",
    "ind": "India",
    "in": "India",
    "filipino": "Philippines",
    "philippine": "Philippines",
    "phl": "Philippines",
    "ph": "Philippines",
    "panamanian": "Panama",
    "pan": "Panama",
    "pa": "Panama",
    "liberian": "Liberia",
    "lbr": "Liberia",
    "lr": "Liberia",
    "marshallese": "Marshall Islands",
    "mhl": "Marshall Islands",
    "british": "United Kingdom",
    "uk": "United Kingdom",
    "gbr": "United Kingdom",
    "russian": "Russian Federation",
    "rus": "Russian Federation",
    "ukrainian": "Ukraine",
    "ukr": "Ukraine",
    "indonesian": "Indonesia",
    "idn": "Indonesia",
    "chinese": "China",
    "chn": "China",
    "vietnamese": "Vietnam",
    "vnm": "Vietnam",
    "georgian": "Georgia",
    "geo": "Georgia",
    "turkish": "Turkey",
    "tur": "Turkey",
    "bangladeshi": "Bangladesh",
    "bgd": "Bangladesh",
    "myanmar": "Myanmar",
    "burmese": "Myanmar",
    "burma": "Myanmar",
    "mmr": "Myanmar",
    "mm": "Myanmar",
}

class ExcelPortalRegistry:
    def __init__(self, excel_path: str = EXCEL_PATH):
        self.excel_path = excel_path
        self._entries: Dict[str, PortalEntry] = {}
        self._load_registry()

    def _load_registry(self):
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        # We load from 'All Countries'
        sheet_name = 'All Countries' if 'All Countries' in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return

        # Find header indices
        header = [str(col).strip().lower() if col is not None else "" for col in rows[0]]
        country_idx = header.index("country") if "country" in header else 0
        url_idx = header.index("official verification url") if "official verification url" in header else 1
        type_idx = header.index("verification type") if "verification type" in header else 2
        coc_idx = header.index("coc") if "coc" in header else 3
        cdc_idx = header.index("cdc / seaman book / sid") if "cdc / seaman book / sid" in header else 5
        captcha_idx = header.index("captcha") if "captcha" in header else 7
        auth_idx = header.index("official authority / source") if "official authority / source" in header else 9
        notes_idx = header.index("verification notes") if "verification notes" in header else 10

        for r in rows[1:]:
            country_name = r[country_idx]
            if not country_name:
                continue
            country_str = str(country_name).strip()

            url_val = str(r[url_idx]).strip() if r[url_idx] and str(r[url_idx]).strip() != "None" else None
            ver_type = str(r[type_idx]).strip() if r[type_idx] else "UNKNOWN"
            coc_supported = True if r[coc_idx] and "yes" in str(r[coc_idx]).lower() else False
            cdc_supported = True if r[cdc_idx] and "yes" in str(r[cdc_idx]).lower() else False
            
            # Check if captcha is observed
            has_captcha = False
            if r[captcha_idx]:
                captcha_str = str(r[captcha_idx]).lower()
                if "yes" in captcha_str or "captcha" in captcha_str and "not observed" not in captcha_str:
                    has_captcha = True

            entry = PortalEntry(
                country=country_str,
                url=url_val,
                verification_type=ver_type,
                supports_coc=coc_supported,
                supports_cdc=cdc_supported,
                has_captcha=has_captcha,
                authority=str(r[auth_idx]).strip() if r[auth_idx] else None,
                notes=str(r[notes_idx]).strip() if r[notes_idx] else None
            )

            self._entries[country_str.lower()] = entry

    def lookup(self, nationality_or_country: str) -> Optional[PortalEntry]:
        """Lookup verification portal metadata by nationality or country name."""
        if not nationality_or_country:
            return None
        
        query = nationality_or_country.strip().lower()

        # Check demonym/alias
        if query in NATIONALITY_ALIASES:
            query = NATIONALITY_ALIASES[query].lower()

        # Exact match
        if query in self._entries:
            return self._entries[query]

        # Partial substring match
        for key, entry in self._entries.items():
            if query in key or key in query:
                return entry

        return None

# Singleton instance
registry = ExcelPortalRegistry()
